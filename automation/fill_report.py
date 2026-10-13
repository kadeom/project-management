#!/usr/bin/env python3
"""Вставить скриншоты в отчёт вместо меток «[ Вставьте скриншот lrN_KK: … ]».

    py -m automation.fill_report

Скриншот ищется как automation/results/shots/<id>.png (или .jpg). Если к одному ID
нужно несколько картинок (до/после), назовите их <id>_a.png, <id>_b.png — вставятся
по порядку. Метки без файлов остаются в отчёте, их список печатается в конце.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

MARK = re.compile(r"\[ Вставьте скриншот (lr\d+_\d+): .* \]")


def images_for(shots: Path, sid: str) -> list[Path]:
    found = []
    for ext in ("png", "jpg", "jpeg"):
        found += [shots / f"{sid}.{ext}"]
        found += sorted(shots.glob(f"{sid}_*.{ext}"))
    return [p for p in found if p.exists()]


def fill(src: Path, dst: Path, shots: Path, width_cm: float = 16.0) -> tuple[list[str], list[str]]:
    from docx import Document
    from docx.shared import Cm

    doc = Document(src)
    done, missing = [], []
    for par in doc.paragraphs:
        m = MARK.fullmatch(par.text.strip())
        if not m:
            continue
        sid = m[1]
        imgs = images_for(shots, sid)
        if not imgs:
            missing.append(sid)
            continue
        for r in par.runs:
            r.text = ""
        run = par.runs[0] if par.runs else par.add_run()
        for i, img in enumerate(imgs):
            if i:
                run.add_break()
            run.add_picture(str(img), width=Cm(width_cm))
        done.append(sid)
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dst)
    return done, missing


HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
from pm import paths  # noqa: E402


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--shots", type=Path, default=paths.SHOTS)
    ap.add_argument("--src", type=Path, default=paths.OUT / paths.REPORT_DOCX)
    ap.add_argument("--dst", type=Path, default=paths.RESULTS / paths.REPORT_READY_DOCX)
    a = ap.parse_args()
    done, missing = fill(a.src, a.dst, a.shots)
    print(f"вставлено: {len(done)}; без скриншота: {len(missing)} {missing}")
    print(f"→ {a.dst}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
