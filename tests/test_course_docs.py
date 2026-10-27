"""Проверки генераторов СРС (SIW) и методических указаний."""
from __future__ import annotations

from pathlib import Path

from docx import Document

from pm import md_docx, model, paths, report_siw


def test_markdown_front_matter_and_blocks():
    meta, body = md_docx.load_markdown(paths.GUIDELINES_MD)
    assert meta["title"]
    blocks = md_docx.parse(body)
    kinds = {b[0] for b in blocks}
    assert {"h", "p", "ul", "ol", "table"} <= kinds
    # Продолжение пункта списка не должно становиться отдельным абзацем.
    bullets = [b[1] for b in blocks if b[0] == "ul"]
    assert any("дополнительные задания" in " ".join(items) for items in bullets)


def test_siw_ids_and_shots_are_consistent():
    data = report_siw.load(paths.SIW_YAML)
    tasks = [t for s in data["sections"] for t in s.get("tasks", [])]
    ids = [t["id"] for t in tasks]
    shots = [x for t in tasks for x in t.get("shots", [])]
    assert len(ids) == len(set(ids))
    assert shots == paths.SIW_SHOT_IDS
    assert {s["lab"] for s in data["sections"]} == set(range(1, 10))


def test_siw_report_docx_contains_all_labs(tmp_path: Path):
    project = model.load(paths.PROJECT_YAML)
    data = report_siw.load(paths.SIW_YAML)
    out = report_siw.docx(project, data, tmp_path / "siw-report.docx")
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    for lab in range(1, 10):
        assert f"ЛР{lab}." in text
    assert "Вставьте скриншот siw_lr1_01" in text
    assert len(doc.tables) >= 1


def test_guidelines_docx_is_rendered(tmp_path: Path):
    meta, body = md_docx.load_markdown(paths.GUIDELINES_MD)
    out = md_docx.render_docx(md_docx.parse(body), tmp_path / "guidelines.docx", meta=meta)
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Методические указания" in text
    assert "ЛР9" in text
    assert "Дополнительные задания 10 и 12" in text
    assert "СРС" in text


def test_project_artifacts_docx_is_rendered(tmp_path: Path):
    meta, body = md_docx.load_markdown(paths.ARTIFACTS_MD)
    out = md_docx.render_docx(md_docx.parse(body), tmp_path / "artifacts.docx", meta=meta)
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Устав проекта" in text
    assert "Реестр рисков" in text
    assert "RACI" in text
    assert len(doc.tables) >= 8


def test_course_docs_are_anonymized():
    """В новых материалах не должно быть домена/имён из приватного прототипа."""
    files = [paths.GUIDELINES_MD, paths.SIW_YAML, paths.ARTIFACTS_MD]
    forbidden = ("phish", "kadeom/phish-eye", "Кошевой", "ИБ-1-23", "@gmail.com")
    for path in files:
        text = Path(path).read_text(encoding="utf-8").lower()
        for token in forbidden:
            assert token.lower() not in text, f"{path}: найден {token!r}"
