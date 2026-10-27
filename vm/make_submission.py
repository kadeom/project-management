#!/usr/bin/env python3
"""Собрать папку для сдачи: .mpp под итоговыми именами + отчёты.

    python vm/make_submission.py
    python vm/make_submission.py --force   # перезаписать существующие файлы

Скрипт идемпотентен: без ``--force`` существующие файлы не удаляет и не
перезаписывает. Имена и состав — из ``paths.SUBMISSION_FILES``; это единственный
источник соответствия «рабочий XML → итоговый .mpp».
"""
from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT / "src"))

from pm import paths  # noqa: E402

FILES = paths.SUBMISSION_FILES
REPORT = paths.SUBMISSION_REPORT
SIW = paths.SUBMISSION_SIW


def _source_mpp(xml_name: str) -> Path | None:
    """Найти .mpp для XML: сначала out/, затем automation/results/."""
    stem = Path(xml_name).stem
    for folder in (paths.OUT, paths.RESULTS):
        p = folder / f"{stem}.mpp"
        if p.exists():
            return p
    return None


def rename_in_doc(doc) -> None:
    """Заменить рабочие XML-имена в тексте на итоговые .mpp (длинные — первыми)."""
    pairs = sorted(((Path(src).name, f"{dst}.mpp") for src, dst, *_ in FILES),
                   key=lambda x: -len(x[0]))

    def fix(par):
        for run in par.runs:
            text = run.text
            for old, new in pairs:
                text = text.replace(old, new)
            if text != run.text:
                run.text = text

    for par in doc.paragraphs:
        fix(par)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    fix(par)


def add_file_table(doc) -> None:
    """Добавить таблицу файлов перед разделом «Выводы», если его ещё нет."""
    anchors = [p for p in doc.paragraphs if p.style.name == "Heading 1" and p.text.strip() == "Выводы"]
    if not anchors:
        return
    if any("Файлы проекта" in p.text for p in doc.paragraphs):
        return
    anchor = anchors[0]
    anchor.insert_paragraph_before("Файлы проекта", style="Heading 1")
    anchor.insert_paragraph_before(
        "MS Project хранит один план в одном .mpp, поэтому каждое состояние плана "
        "и каждое дополнительное задание сохранены отдельным файлом:")
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = doc.tables[0].style if doc.tables else "Table Grid"
    for cell, head in zip(tbl.rows[0].cells, ("Файл", "Работа", "Содержание"), strict=True):
        cell.text = head
    for _, dst, lab, what in FILES:
        row = tbl.add_row().cells
        row[0].text, row[1].text, row[2].text = f"{dst}.mpp", lab, what
    anchor._element.addprevious(tbl._element)


def copy(src: Path, dst: Path, force: bool = False) -> None:
    if dst.exists() and not force:
        print(f"есть, пропускаю: {dst.name}")
        return
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dst)
    print(f"скопировано: {dst.name}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--force", action="store_true", help="перезаписать существующие файлы")
    args = ap.parse_args()

    out = paths.SUBMISSION_DIR
    out.mkdir(parents=True, exist_ok=True)

    missing: list[str] = []
    for src, dst, *_ in FILES:
        mpp = _source_mpp(src)
        if mpp is None:
            missing.append(f"{Path(src).stem}.mpp")
            continue
        copy(mpp, out / f"{dst}.mpp", args.force)
    if missing:
        print(f"нет .mpp (соберите scripts/xml_to_mpp.py): {', '.join(missing)}", file=sys.stderr)

    # Отчёт со скриншотами: results/ готовый приоритетнее шаблона out/.
    report_src = paths.RESULTS / paths.REPORT_READY_DOCX
    if not report_src.exists():
        report_src = paths.OUT / paths.REPORT_DOCX
    report_dst = out / REPORT
    if report_src.exists():
        if report_dst.exists() and not args.force:
            print(f"есть, пропускаю: {REPORT}")
        else:
            from docx import Document
            doc = Document(report_src)
            rename_in_doc(doc)
            add_file_table(doc)
            doc.save(report_dst)
            print(f"собран отчёт: {REPORT}")
    else:
        print(f"нет отчёта {paths.REPORT_DOCX} — сначала build.py", file=sys.stderr)

    siw_src = paths.RESULTS / paths.SIW_REPORT_READY_DOCX
    if not siw_src.exists():
        siw_src = paths.OUT / paths.SIW_REPORT_DOCX
    if siw_src.exists():
        copy(siw_src, out / SIW, args.force)
    else:
        print(f"нет отчёта {paths.SIW_REPORT_DOCX} — сначала build_siw.py", file=sys.stderr)

    print(f"→ {out}")
    return 0 if not missing else 1


if __name__ == "__main__":
    raise SystemExit(main())
